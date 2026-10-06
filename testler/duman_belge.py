# -*- coding: utf-8 -*-
"""BELGE KUTUPHANESI (06.10.2026, Asama 2 / E77). Jeton yakmaz.

Patron PDF / Word / Excel / CSV / metin yukler; belge parcalara bolunur,
her parca YERINI tasir; ajanlar (`belge_ara`, brifing) ve lider (triyaj)
ilgili parcayi KAYNAGIYLA gorur.

Sinanan: bes tur gercek dosyadan metin ve yer (s.2, bolum, satir) cikar ·
arama dogru parcayi ve kaynagini getirir, ilgisizi getirmez · desteklenmeyen
tur / metinsiz PDF / ayni belge reddedilir (sessiz "yuklendi" yok) · brifing
ve triyaj ilgili parcayi kaynagiyla tasir · gercek MCP sureci `belge_ara`
cevabi · kaldirilan belge aranmaz · panel ve kabuk.
"""

from __future__ import annotations

import _temizlik  # noqa: F401 -- gecici klasorleri siler; ILK ice aktarma

import base64
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

KOK = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(KOK))

from ofis import Ofis, Veritabani, belge as B  # noqa: E402
from ofis.motor.sahte import SahteMotor  # noqa: E402
from ofis.sohbet import Sohbet  # noqa: E402
from _pdf import pdf_yap  # noqa: E402

GECEN = KALAN = 0


def sina(baslik: str, kosul: bool, not_: str = "") -> None:
    global GECEN, KALAN
    if kosul:
        GECEN += 1
        print(f"  [gecti] {baslik}")
    else:
        KALAN += 1
        print(f"  [KALDI] {baslik} {not_}")


def hata(islev, *a, **k) -> str:
    try:
        islev(*a, **k)
    except B.BelgeHatasi as exc:
        return str(exc)
    return ""


d = Path(tempfile.mkdtemp(prefix="ofis-belge-"))
db = Veritabani(d / "o.sqlite3")
ofis = Ofis.proje_ac(db, "Belge", amac="a", basari_olcutu="b", kok_klasor=str(d))
kn = ofis.kanat_ac("GENEL", "GE")
ofis.kadro.ise_al("lale", "lider", kanat_id=kn, lider_mi=True, motor="sahte", klasor=str(d))
ofis.kadro.ise_al("cem", "backend", kanat_id=kn, motor="sahte", klasor=str(d))

print("== belge kutuphanesi ==\n")
print("1. bes tur")
pdf = pdf_yap(["Semu Parfum magaza kurallari", "Iade suresi 30 gundur. Kargo ucreti alicidan alinir."])
r = B.yukle(ofis, "politika.pdf", pdf)
sina("PDF yuklendi ve parcalandi", r["parca"] >= 1 and (d / ".ofis/belgeler/politika.pdf").is_file(), str(r))
import docx  # noqa: E402
w = docx.Document()
w.add_heading("Marka tonu", level=1)
w.add_paragraph("Musteriye her zaman siz diye hitap edilir; emoji kullanilmaz.")
w.add_heading("Teslimat", level=1)
w.add_paragraph("Siparisler 2 is gununde kargoya verilir.")
w.save(str(d / "rehber.docx"))
r = B.yukle_yoldan(ofis, d / "rehber.docx")
sina("Word yuklendi", r["parca"] >= 1)
import openpyxl  # noqa: E402
kt = openpyxl.Workbook()
sy = kt.active
sy.title = "Fiyatlar"
for satir in (("Urun", "Fiyat"), ("Belle", 1490), ("Brave", 1650)):
    sy.append(satir)
kt.save(str(d / "fiyat.xlsx"))
sina("Excel yuklendi", B.yukle_yoldan(ofis, d / "fiyat.xlsx")["parca"] >= 1)
sina("CSV yuklendi", B.yukle(ofis, "stok.csv", "urun,adet\nBelle,12\nBrave,0\n".encode())["parca"] >= 1)
sina("Markdown yuklendi", B.yukle(ofis, "notlar.md",
                                   "# Kampanya\n\nKasim indirimi yuzde 15.\n".encode())["parca"] >= 1)

print("\n2. ret (sessiz 'yuklendi' yok)")
sina("desteklenmeyen tur reddedildi", "desteklenmiyor" in hata(B.yukle, ofis, "sunum.pptx", b"x"))
sina("metinsiz PDF reddedildi", "METIN yok" in hata(B.yukle, ofis, "bos.pdf", pdf_yap([""])),
     hata(B.yukle, ofis, "bos2.pdf", pdf_yap([" "])))
sina("ayni belge ikinci kez reddedildi", "zaten yuklu" in hata(B.yukle, ofis, "kopya.pdf", pdf))
sina("bos dosya reddedildi", bool(hata(B.yukle, ofis, "x.txt", b"")))
sina("reddedilen dosya diske birakilmadi",
     not (d / ".ofis/belgeler/bos.pdf").exists() and not (d / ".ofis/belgeler/sunum.pptx").exists())

print("\n3. arama ve kaynak")
s = B.ara(db, ofis.proje_id, "iade süresi kaç gün")
sina("iade sorusu PDF'in 2. sayfasini getirdi", s and s[0]["belge"] == "politika.pdf"
     and s[0]["yer"] == "s.2" and "30" in s[0]["metin"], str(s[:1]))
s = B.ara(db, ofis.proje_id, "musteriye nasil hitap edilir")
sina("Word: dogru bolum", s and s[0]["belge"] == "rehber.docx" and "Marka tonu" in s[0]["yer"], str(s[:1]))
s = B.ara(db, ofis.proje_id, "Brave fiyati")
sina("Excel: sayfa ve satir yeri", s and s[0]["belge"] == "fiyat.xlsx" and "Fiyatlar satir" in s[0]["yer"],
     str(s[:1]))
sina("ilgisiz sorgu bos doner (uydurma parca yok)",
     B.ara(db, ofis.proje_id, "uzay roketi yakit") == [])

print("\n4. brifing ve triyaj")
g = ofis.tahta.gorev_ac("Iade sorusuna cevap", neden="musteri iade suresini soruyor",
                        ana_hedef="musteriye iade suresini yaz", kabul_olcutu="cevap dogru",
                        kanat_id=kn)
k2 = ofis.brifing.k2_kunye(g["no"])
sina("brifing ilgili parcayi KAYNAGIYLA tasiyor",
     "Patronun belgeleri" in k2 and "[politika.pdf · s.2]" in k2 and "KAYNAGINI" in k2, k2[-600:])
tri = Sohbet(ofis, motor_fabrikasi=lambda a, **k: SahteMotor()).triyaj_metni("iade suresi kac gun?")
sina("triyaj (liderin on kapisi) belgeyi kaynagiyla goruyor",
     "PATRONUN BELGELERI" in tri and "[politika.pdf · s.2]" in tri)

print("\n5. gercek MCP sureci")
cevre = os.environ.copy()
cevre.update({"OFIS_DB": str(db.yol), "OFIS_PROJE": ofis.proje_id, "OFIS_AJAN": "cem",
              "OFIS_GOREV": str(g["no"])})
p = subprocess.Popen([sys.executable, str(KOK / "ofis" / "mcp_sunucu.py")], stdin=subprocess.PIPE,
                     stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, encoding="utf-8",
                     env=cevre, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
istek = [{"jsonrpc": "2.0", "id": 1, "method": "initialize",
          "params": {"protocolVersion": "2025-06-18", "capabilities": {}}},
         {"jsonrpc": "2.0", "id": 2, "method": "tools/list"},
         {"jsonrpc": "2.0", "id": 3, "method": "tools/call",
          "params": {"name": "belge_ara", "arguments": {"sorgu": "kargo ucreti kimden"}}}]
cikti, _ = p.communicate("\n".join(json.dumps(x) for x in istek) + "\n", timeout=60)
cev = {json.loads(x)["id"]: json.loads(x) for x in cikti.splitlines() if x.strip().startswith("{")}
adlar = [t["name"] for t in cev.get(2, {}).get("result", {}).get("tools", [])]
sina("belge_ara araci listede", "belge_ara" in adlar)
metin = cev.get(3, {}).get("result", {}).get("content", [{}])[0].get("text", "")
sina("belge_ara kaynakli parca dondurdu", "[politika.pdf · s.2]" in metin and "alicidan" in metin, metin[:200])

print("\n6. kaldir, panel, kabuk")
bid = next(b["id"] for b in B.liste(db, ofis.proje_id) if b["ad"] == "notlar.md")
B.kaldir(ofis, bid[:8])
sina("kaldirilan belge aranmiyor (dosya duruyor)",
     not B.ara(db, ofis.proje_id, "kasim indirimi") and (d / ".ofis/belgeler/notlar.md").exists())
sina("yalniz patron kaldirir", "patron" in hata(B.kaldir, ofis, bid, kim="cem"))
from fastapi.testclient import TestClient  # noqa: E402
from ofis.panel import olustur  # noqa: E402
ist = TestClient(olustur(db.yol))
sina("panel belge listesi", len(ist.get("/api/belgeler").json()["belgeler"]) == 4)
r = ist.post("/api/belge", json={"ad": "sss.txt", "veri": base64.b64encode(
    "Sikca sorulan: hediye paketi ucretsizdir.".encode()).decode()})
sina("panelden yukleme", r.status_code == 200 and r.json()["parca"] == 1, r.text[:120])
r = ist.post("/api/belge", json={"ad": "x.exe", "veri": base64.b64encode(b"MZ").decode()})
sina("panel desteklenmeyen turu 400 ile reddediyor", r.status_code == 400)
sina("panel arama", ist.get("/api/belge/ara", params={"q": "hediye paketi"}).json()["sonuc"][0]["belge"] == "sss.txt")
k = subprocess.run([sys.executable, "-m", "ofis.kabuk", "--db", str(db.yol), "belge", "liste"],
                   capture_output=True, text=True, encoding="utf-8", errors="replace", cwd=str(KOK),
                   creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
sina("kabuk `belge liste`", "politika.pdf" in k.stdout and "sss.txt" in k.stdout, (k.stdout + k.stderr)[-300:])

print(f"\n== sonuc: {GECEN} gecti, {KALAN} kaldi ==")
sys.exit(1 if KALAN else 0)
